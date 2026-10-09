import sqlite3
import json
import os
import openpyxl

print("=== INSPECTING vignan_academic.db ===")
if os.path.exists("vignan_academic.db"):
    conn = sqlite3.connect("vignan_academic.db")
    cursor = conn.cursor()
    tables = cursor.execute("SELECT name FROM sqlite_master WHERE type='table';").fetchall()
    print("Tables:", [t[0] for t in tables])
    for t in tables:
        count = cursor.execute(f"SELECT count(*) FROM {t[0]}").fetchone()[0]
        cols = [c[1] for c in cursor.execute(f"PRAGMA table_info({t[0]})").fetchall()]
        print(f"  {t[0]} ({count} rows): {cols[:6]}...")
    conn.close()
else:
    print("vignan_academic.db not found!")

print("\n=== INSPECTING sample_data JSONs ===")
for fname in ["existing_api_students.json", "existing_teacher_students.json", "backend_subjects.json", "backend_degree_requirements.json", "backend_policies_attendance.json"]:
    fpath = os.path.join("sample_data", fname)
    if os.path.exists(fpath):
        with open(fpath, "r", encoding="utf-8") as f:
            data = json.load(f)
            if isinstance(data, list):
                print(f"{fname}: list with {len(data)} items")
                if len(data) > 0:
                    print("  Sample item keys/data:", list(data[0].keys()) if isinstance(data[0], dict) else data[0])
                    print("  Sample item:", data[0])
            elif isinstance(data, dict):
                print(f"{fname}: dict with {len(data)} keys: {list(data.keys())[:8]}")
                # print first entry
                first_k = list(data.keys())[0]
                print(f"  Sample [{first_k}]:", data[first_k])

print("\n=== INSPECTING exports/ EXCEL FILES ===")
if os.path.exists("exports"):
    for fname in os.listdir("exports"):
        if fname.endswith(".xlsx"):
            fpath = os.path.join("exports", fname)
            wb = openpyxl.load_workbook(fpath, read_only=True)
            print(f"{fname}: sheets = {wb.sheetnames}")
            for sheetname in wb.sheetnames:
                ws = wb[sheetname]
                rows = list(ws.iter_rows(values_only=True, max_row=5))
                print(f"  [{sheetname}] Header / sample ({len(rows)} rows preview):")
                for r in rows[:3]:
                    print("   ", r)
