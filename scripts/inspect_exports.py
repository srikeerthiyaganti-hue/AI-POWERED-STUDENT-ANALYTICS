import openpyxl, os

for f in os.listdir("exports"):
    if f.endswith(".xlsx"):
        path = os.path.join("exports", f)
        wb = openpyxl.load_workbook(path, data_only=True)
        print(f"File: {f}")
        for s in wb.sheetnames:
            ws = wb[s]
            print(f"  Sheet '{s}': {ws.max_row} rows, {ws.max_column} cols")
            headers = [ws.cell(1, c).value for c in range(1, min(ws.max_column + 1, 15))]
            print(f"    Headers: {headers}")
