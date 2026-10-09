import openpyxl

wb1 = openpyxl.load_workbook("exports/vignan_existing_students.xlsx", data_only=True)
ws1 = wb1["Students"]
print("=== vignan_existing_students.xlsx [Students] ===")
for row in ws1.iter_rows(values_only=True):
    print(row)

wb2 = openpyxl.load_workbook("exports/vignan_existing_data.xlsx", data_only=True)
ws2 = wb2["Candidate_Students"]
print("\n=== vignan_existing_data.xlsx [Candidate_Students] (First 10) ===")
for i, row in enumerate(ws2.iter_rows(values_only=True)):
    if i < 10:
        print(row)
