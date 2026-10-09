import io
import re
import logging
from typing import List, Dict, Any
from openpyxl import load_workbook
from app.normalizer import clean_text, normalize_course_code, infer_course_category

logger = logging.getLogger("academic_pipeline.excel_extractor")

class AcademicExcelExtractor:
    """Extracts curriculum structures and courses from Excel workbooks."""

    def __init__(self, excel_bytes: bytes, source_document_id: int):
        self.wb = load_workbook(io.BytesIO(excel_bytes), data_only=True)
        self.source_document_id = source_document_id

    def extract_courses(self) -> Dict[int, List[Dict[str, Any]]]:
        """
        Parses sheets in workbook. Detects semester from sheet name (e.g. 'Sem 1', 'Semester 2')
        or from a 'Semester' column.
        Returns dict mapping semester_number -> list of courses.
        """
        semesters_data: Dict[int, List[Dict[str, Any]]] = {}

        for sheet_name in self.wb.sheetnames:
            ws = self.wb[sheet_name]
            
            # Infer semester from sheet name if possible
            sem_m = re.search(r'(\d+)', sheet_name)
            sheet_sem = int(sem_m.group(1)) if sem_m else 1

            rows = list(ws.iter_rows(values_only=True))
            if not rows or len(rows) < 2:
                continue

            # Find header row
            header_idx = -1
            col_map = {}
            for r_idx, row in enumerate(rows[:5]):
                row_str = [str(c).lower().strip() if c is not None else "" for c in row]
                for c_idx, val in enumerate(row_str):
                    if "code" in val: col_map["code"] = c_idx
                    elif "name" in val or "title" in val or "course" in val: col_map["name"] = c_idx
                    elif val in ["credits", "credit", "c"]: col_map["credits"] = c_idx
                    elif val in ["lecture", "l"]: col_map["l"] = c_idx
                    elif val in ["tutorial", "t"]: col_map["t"] = c_idx
                    elif val in ["practical", "lab", "p"]: col_map["p"] = c_idx
                    elif "prereq" in val: col_map["prereq"] = c_idx
                    elif "sem" in val: col_map["sem"] = c_idx

                if "code" in col_map and "name" in col_map:
                    header_idx = r_idx
                    break

            if header_idx == -1:
                # Default mapping if no matching header row found
                col_map = {"code": 0, "name": 1, "credits": 2}
                header_idx = 0

            # Parse data rows
            for row in rows[header_idx + 1:]:
                if not row or not any(row):
                    continue

                code_val = row[col_map["code"]] if "code" in col_map and col_map["code"] < len(row) else None
                name_val = row[col_map["name"]] if "name" in col_map and col_map["name"] < len(row) else None

                if not code_val or not name_val:
                    continue

                code = normalize_course_code(str(code_val))
                name = clean_text(str(name_val))

                if len(code) < 3 or len(name) < 2:
                    continue

                credits = 3.0
                if "credits" in col_map and col_map["credits"] < len(row):
                    try:
                        credits = float(row[col_map["credits"]])
                    except (ValueError, TypeError):
                        pass

                l = float(row[col_map["l"]]) if "l" in col_map and col_map["l"] < len(row) and str(row[col_map["l"]]).replace('.','').isdigit() else 3.0
                t = float(row[col_map["t"]]) if "t" in col_map and col_map["t"] < len(row) and str(row[col_map["t"]]).replace('.','').isdigit() else 0.0
                p = float(row[col_map["p"]]) if "p" in col_map and col_map["p"] < len(row) and str(row[col_map["p"]]).replace('.','').isdigit() else 0.0

                prereq = clean_text(str(row[col_map["prereq"]])) if "prereq" in col_map and col_map["prereq"] < len(row) and row[col_map["prereq"]] else None

                row_sem = sheet_sem
                if "sem" in col_map and col_map["sem"] < len(row):
                    try:
                        row_sem = int(row[col_map["sem"]])
                    except (ValueError, TypeError):
                        pass

                course_item = {
                    "course_code": code,
                    "course_name": name,
                    "credits": credits,
                    "lecture_hours": l,
                    "tutorial_hours": t,
                    "practical_hours": p,
                    "prerequisites": prereq,
                    "course_category": infer_course_category(code, name),
                    "source_document_id": self.source_document_id
                }

                semesters_data.setdefault(row_sem, []).append(course_item)

        return semesters_data
